import base64
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, timedelta

from odoo import _, api, fields, models
from odoo.exceptions import UserError

from odoo.addons.odoo_rrhh_custom.models.employee_overtime import JOURNEY_TYPES

# Valor de la hora = Salario LPS / 194.85 (jornada mensual efectiva).
MONTHLY_HOURS = 194.85

# Tipo de jornada cuando un día no se puede clasificar (salida ilegible) o el
# reporte no trae detalle diario.
FALLBACK_JOURNEY_TYPE = 'diurna'

SPANISH_MONTHS = {
    'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6,
    'jul': 7, 'ago': 8, 'sep': 9, 'set': 9, 'oct': 10, 'nov': 11, 'dic': 12,
}

_DIGITS_RE = re.compile(r'^\d+$')

# "... 17 Ago 2026 al 23 Ago 2026"
PERIOD_RE = re.compile(
    r'(\d{1,2})\s+([A-Za-zÀ-ſ]{3,})\.?\s+(\d{4})\s*al\s*'
    r'(\d{1,2})\s+([A-Za-zÀ-ſ]{3,})\.?\s+(\d{4})',
    re.IGNORECASE,
)
# Encabezado de columna de día: "Lun 17 Ago"
DAY_HEADER_RE = re.compile(r'^(lun|mar|mie|jue|vie|sab|dom)\s+(\d{1,2})\s+([a-z]{3,})')


def _norm(value):
    """minúsculas, sin acentos, sin saltos de línea, espacios colapsados."""
    text = (value if isinstance(value, str) else '' if value is None else str(value))
    text = text.replace('\n', ' ').replace('\r', ' ').strip().lower()
    text = ''.join(c for c in unicodedata.normalize('NFKD', text)
                   if not unicodedata.combining(c))
    return ' '.join(text.split())


def _parse_float(value):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(',', '.')
    if not text or text in ('-', '?'):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _parse_clock(value):
    """Hora del reloj -> float de horas [0, 48). None si no se puede leer."""
    if value in (None, '', '-', '?'):
        return None
    if isinstance(value, (int, float)):
        v = float(value)
        if 0 <= v < 2:          # fracción de día (xlsx/xls serial)
            return v * 24
        if 2 <= v < 48:         # ya viene en horas
            return v
        return None
    m = re.match(r'^(\d{1,2}):(\d{2})', str(value).strip())
    if m:
        h, mi = int(m.group(1)), int(m.group(2))
        if h < 48 and mi < 60:
            return h + mi / 60
    return None


def _clock_id(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip() if value not in (None, '') else ''


def _mk_date(day, month_name, year):
    month = SPANISH_MONTHS.get(_norm(month_name)[:3])
    if not month:
        return None
    try:
        return date(int(year), month, int(day))
    except (ValueError, TypeError):
        return None


class EmployeeOvertimeBulkImport(models.TransientModel):
    _name = 'employee.overtime.bulk.import'
    _description = 'Carga masiva de horas extra (reporte de asistencia)'

    state = fields.Selection([
        ('choose', 'Subir archivo'),
        ('done', 'Resultado'),
    ], default='choose', required=True)
    file = fields.Binary(string='Reporte de asistencia (.xls / .xlsx)')
    filename = fields.Char(string='Nombre del archivo')
    week_start_override = fields.Date(
        string='Fecha de la semana (opcional)',
        help='Solo si el archivo no trae el rango "… al …" en el encabezado.')
    replace_existing = fields.Boolean(
        string='Reemplazar cargas previas de esta semana',
        help='Si ya se importó el reporte de esta semana, borra esas horas extra '
             '(solo las de origen "Reporte de asistencia", misma fecha) y las '
             'vuelve a crear. Si se deja desmarcado, esos empleados se omiten.')
    result_summary = fields.Text(string='Resultado', readonly=True)
    created_count = fields.Integer(string='Registros creados', readonly=True)
    no_extra_count = fields.Integer(string='Sin horas extra', readonly=True)
    duplicate_count = fields.Integer(string='Ya cargadas (omitidas)', readonly=True)
    skipped_count = fields.Integer(string='Omitidas por error', readonly=True)
    created_overtime_ids = fields.Many2many(
        'employee.overtime', string='Horas extra creadas')

    # ------------------------------------------------------------------ #
    #  Acciones
    # ------------------------------------------------------------------ #
    def action_import(self):
        self.ensure_one()
        self._check_especialista()
        if not self.file:
            raise UserError(_('Selecciona el reporte de asistencia antes de importar.'))

        res = self._parse()
        errors, warnings = res['errors'], res['warnings']

        created = self.env['employee.overtime']
        if res['vals_list']:
            created = self.env['employee.overtime'].with_context(
                overtime_bulk_import=True).create(res['vals_list'])

        summary = [
            _('Semana del reporte: %(a)s a %(b)s',
              a=res['period_start'], b=res['period_end']),
            _('Empleados con horas extra: %s', res['employees_ok']),
            _('Registros creados (uno por tipo de jornada): %s', len(created)),
            _('Empleados sin horas extra (omitidos): %s', res['no_extra']),
        ]
        if res['scaled']:
            summary.append(_('Empleados cuyo "Total Horas" se repartió por tipo '
                             'según el patrón diario (la suma diaria no cuadraba '
                             'exacto): %s', res['scaled']))
        if res['holidays_worked']:
            summary.append(_('Feriados/domingos trabajados registrados (pagan '
                             'doble ese día en la planilla): %s',
                             res['holidays_worked']))
        if res['replaced']:
            summary.append(_('Cargas previas de esta semana borradas y '
                             'reemplazadas: %s', res['replaced']))
        if res['duplicates']:
            summary.append(_('Empleados ya cargados de esta semana (omitidos, '
                             'marca "Reemplazar" para rehacerlos): %s',
                             res['duplicates']))
        if warnings:
            summary.append('')
            summary.append(_('Avisos:'))
            summary.extend(warnings)
        if errors:
            summary.append('')
            summary.append(_('Filas omitidas por error: %s', len(errors)))
            summary.extend(errors)
        elif not warnings:
            summary.append(_('Sin errores.'))

        self.write({
            'state': 'done',
            'result_summary': '\n'.join(str(s) for s in summary),
            'created_count': len(created),
            'no_extra_count': res['no_extra'],
            'duplicate_count': res['duplicates'],
            'skipped_count': len(errors),
            'created_overtime_ids': [(6, 0, created.ids)],
        })
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def action_open_records(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Horas extra creadas'),
            'res_model': 'employee.overtime',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.created_overtime_ids.ids)],
        }

    # ------------------------------------------------------------------ #
    #  Parseo del reporte
    # ------------------------------------------------------------------ #
    def _check_especialista(self):
        group = self.env.ref('odoo_rrhh_custom.group_rrhh_especialista',
                             raise_if_not_found=False)
        if not group or group not in self.env.user.groups_id:
            raise UserError(_('Solo un Especialista RRHH puede hacer la carga '
                              'masiva de horas extra.'))

    def _parse(self):
        matrix = self._read_matrix()
        if not matrix:
            raise UserError(_('El archivo está vacío.'))

        period_start, period_end = self._extract_period(matrix)
        header_idx, cols = self._locate_columns(matrix, period_start)
        data_start = self._first_data_row(matrix, header_idx, cols)
        holiday_set = self._load_holidays(period_start, period_end)

        previous = self.env['employee.overtime'].search([
            ('source', '=', 'attendance_import'),
            ('date', '=', period_start),
        ])
        replaced = 0
        if previous and self.replace_existing:
            replaced = len(previous)
            previous.unlink()
            previous = self.env['employee.overtime']
        already_loaded_ids = set(previous.mapped('employee_id').ids)

        vals_list, errors, warnings = [], [], []
        no_extra = duplicates = employees_ok = scaled = 0
        holidays_worked = 0
        for r in range(data_start, len(matrix)):
            row = matrix[r]
            clock_id = _clock_id(row[cols['id']] if cols['id'] < len(row) else '')
            name = str(row[cols['name']] or '').strip() if cols['name'] < len(row) else ''
            # Solo son datos las filas cuyo 'ID' es un número (así se ignoran los
            # bloques de encabezado repetidos a media hoja y las filas vacías).
            if not _DIGITS_RE.match(clock_id):
                continue
            tag = 'ID %s (%s)' % (clock_id, name or _('sin nombre'))

            # Feriado/domingo trabajado (paga doble el día): se registra sin
            # importar si ese día generó o no horas extra.
            holidays_worked += self._register_holidays_worked(
                row, cols, holiday_set, name)

            report_total = (_parse_float(row[cols['value']])
                            if cols['value'] < len(row) else None)
            has_any_hours = report_total and report_total > 0 or any(
                (_parse_float(row[b['horas']]) or 0) > 0
                for b in (cols.get('day_blocks') or []) if b['horas'] < len(row))
            if not has_any_hours:
                no_extra += 1
                continue

            employee = self._match_employee(name)
            if not employee:
                errors.append(_('%s: no se encontró un empleado único con ese '
                                'nombre en Odoo.', tag))
                continue
            if employee.id in already_loaded_ids:
                duplicates += 1
                continue
            hourly_rate = self._resolve_hourly_rate(employee)
            if not hourly_rate:
                errors.append(_('%s: el empleado no tiene salario activo, no se '
                                'puede calcular el monto.', tag))
                continue

            company = employee.company_id or self.env.company
            raw_buckets, day_note = self._classify_row(row, cols, holiday_set, company)
            raw_total = sum(raw_buckets.values())

            # "Total Horas" es la cantidad real de horas extra (confirmado por el
            # cliente). El detalle diario solo decide el REPARTO por tipo de
            # jornada; se escala para que la suma cuadre con "Total Horas".
            if report_total and report_total > 0:
                if raw_total > 0:
                    k = report_total / raw_total
                    buckets = {jt: h * k for jt, h in raw_buckets.items()}
                    if abs(raw_total - report_total) > 0.01:
                        scaled += 1
                    target = round(report_total, 2)
                else:
                    buckets = {FALLBACK_JOURNEY_TYPE: report_total}
                    target = round(report_total, 2)
                    day_note = _('sin detalle diario; todo como %s',
                                 dict(JOURNEY_TYPES)[FALLBACK_JOURNEY_TYPE])
            elif cols['mode'] != 'total' and raw_total > 0:
                # Formato viejo (sin columna "Total Horas"): se usa la suma diaria.
                buckets = dict(raw_buckets)
                target = round(raw_total, 2)
                day_note = (day_note + '; ' if day_note else '') + _(
                    'sin "Total Horas"; se usó la suma diaria')
            else:
                # Modo "Total Horas" y la celda de total está vacía -> sin horas extra.
                no_extra += 1
                continue

            # El desglose por tipo debe cuadrar EXACTO con el total: se redondea
            # cada parte y el residuo del redondeo se carga a la parte mayor.
            buckets = {jt: round(h, 2) for jt, h in buckets.items() if round(h, 2) > 0}
            residual = round(target - round(sum(buckets.values()), 2), 2)
            if residual and buckets:
                biggest = max(buckets, key=buckets.get)
                buckets[biggest] = round(buckets[biggest] + residual, 2)

            parts = ', '.join('%s %.2f h' % (dict(JOURNEY_TYPES)[jt], h)
                              for jt, h in buckets.items())
            base_just = _('Horas extra semana %(a)s a %(b)s (reporte de '
                          'asistencia). Desglose: %(parts)s.',
                          a=period_start, b=period_end, parts=parts)
            if day_note:
                base_just += ' [%s]' % day_note

            for journey_type, hrs in buckets.items():
                if hrs <= 0:
                    continue
                vals_list.append({
                    'employee_id': employee.id,
                    'company_id': employee.company_id.id,
                    'date': period_start,
                    'hours': round(hrs, 2),
                    'hourly_rate': hourly_rate,
                    'journey_type': journey_type,
                    'justification': base_just,
                    'state': 'approved',
                    'source': 'attendance_import',
                    'rrhh_approver_id': self._rrhh_employee(employee.company_id).id,
                })
            employees_ok += 1

        if not vals_list and not no_extra and not errors and not duplicates:
            raise UserError(_('No se encontraron filas de empleados en el reporte. '
                              '¿El formato del archivo es el esperado?'))
        return {
            'vals_list': vals_list,
            'no_extra': no_extra,
            'duplicates': duplicates,
            'replaced': replaced,
            'employees_ok': employees_ok,
            'scaled': scaled,
            'holidays_worked': holidays_worked,
            'errors': errors,
            'warnings': warnings,
            'period_start': period_start,
            'period_end': period_end,
        }

    def _register_holidays_worked(self, row, cols, holiday_set, name):
        """Si algún día de la fila cae en domingo/feriado y el empleado marcó
        entrada (trabajó), registra payroll.holiday.worked (paga doble ese
        día en la planilla). Devuelve cuántos se crearon nuevos."""
        blocks = cols.get('day_blocks') or []
        if not blocks or not name:
            return 0
        created = 0
        employee = None
        for block in blocks:
            day_date = block['date']
            if not day_date:
                continue
            is_holiday = (day_date.weekday() == 6
                         or any(d == day_date for d, _c in holiday_set))
            if not is_holiday:
                continue
            entrada = row[block['entrada']] if block['entrada'] < len(row) else None
            if _parse_clock(entrada) is None:
                continue  # sin marca de entrada -> no trabajó ese día
            if employee is None:
                employee = self._match_employee(name)
                if not employee:
                    return 0
            exists = self.env['payroll.holiday.worked'].search_count([
                ('employee_id', '=', employee.id), ('date', '=', day_date)])
            if exists:
                continue
            self.env['payroll.holiday.worked'].create({
                'employee_id': employee.id,
                'company_id': employee.company_id.id,
                'date': day_date,
                'source': 'attendance_import',
            })
            created += 1
        return created

    def _classify_row(self, row, cols, holiday_set, company):
        """Devuelve ({journey_type: horas}, nota) para una fila de empleado.

        Cada día se reparte entre tipos de jornada según en qué banda horaria
        caen sus horas extra (contadas desde el fin de jornada). La cantidad
        real la pone el reporte ("Horas" del día); las bandas solo dan la
        proporción.
        """
        buckets = defaultdict(float)
        note = ''
        blocks = cols.get('day_blocks') or []
        if not blocks:
            return buckets, note
        unreadable_days = 0
        for block in blocks:
            hrs = (_parse_float(row[block['horas']])
                   if block['horas'] < len(row) else None)
            if not hrs or hrs <= 0:
                continue
            entrada = row[block['entrada']] if block['entrada'] < len(row) else None
            salida = row[block['salida']] if block['salida'] < len(row) else None
            weights = self._classify_day(
                block['date'], entrada, salida, company, holiday_set)
            if _parse_clock(salida) is None and 'feriado' not in weights:
                unreadable_days += 1
            total_w = sum(weights.values()) or 1.0
            for journey_type, w in weights.items():
                buckets[journey_type] += hrs * (w / total_w)
        if unreadable_days:
            note = _('%s día(s) sin salida legible → %s', unreadable_days,
                     dict(JOURNEY_TYPES)[FALLBACK_JOURNEY_TYPE])
        return buckets, note

    def _classify_day(self, day_date, entrada, salida, company, holiday_set):
        """Reparto de las horas extra de un día como pesos {journey_type: horas
        de la ventana}. Domingo/feriado -> todo 'feriado'. Salida ilegible ->
        todo 'diurna'."""
        if day_date and (day_date.weekday() == 6
                         or (day_date, company.id) in holiday_set
                         or (day_date, False) in holiday_set):
            return {'feriado': 1.0}
        s = _parse_clock(salida)
        if s is None:
            return {FALLBACK_JOURNEY_TYPE: 1.0}
        e = _parse_clock(entrada)
        o_start = company.payroll_ordinary_start or 7.0
        if day_date and day_date.weekday() == 5:            # sábado, jornada corta
            o_end = company.payroll_saturday_end or 12.0
        else:
            o_end = company.payroll_ordinary_end or 17.0
        span = (o_end - o_start) if o_end > o_start else 10.0
        shift_start = max(e, o_start) if e is not None else o_start
        shift_end = shift_start + span
        if s < shift_start:                                 # cruzó medianoche
            s += 24
        if s <= shift_end + 0.01:
            # sin ventana tras la jornada; ¿entró antes de tiempo (H.E. matutina)?
            if e is not None and e < o_start - 0.01:
                return {'nocturna' if e < 5.0 else 'diurna': 1.0}
            return {FALLBACK_JOURNEY_TYPE: 1.0}
        return self._split_ot_bands(shift_end, s, company)

    def _split_ot_bands(self, start, end, company):
        """Reparte la ventana [start, end] (horas del reloj, puede pasar de 24)
        en las bandas Diurna / Mixta / Nocturna / Prolongación."""
        mixta = company.payroll_ot_mixta_from or 19.0
        noct = company.payroll_ot_nocturna_from or 22.0
        prol = company.payroll_ot_prolongacion_from or 5.0
        prol_abs = prol + 24 if prol < noct else prol       # 05:00 del día siguiente
        edges = [
            (-1e9, mixta, 'diurna'),
            (mixta, noct, 'mixta'),
            (noct, prol_abs, 'nocturna'),
            (prol_abs, 1e9, 'prolongacion'),
        ]
        out = {}
        for lo, hi, name in edges:
            overlap = max(0.0, min(end, hi) - max(start, lo))
            if overlap > 0.0001:
                out[name] = out.get(name, 0.0) + overlap
        return out or {'diurna': 1.0}

    def _load_holidays(self, period_start, period_end):
        recs = self.env['payroll.holiday'].sudo().search([
            ('date', '>=', period_start), ('date', '<=', period_end)])
        return {(r.date, r.company_id.id) for r in recs}

    def _read_matrix(self):
        data = base64.b64decode(self.file)
        name = (self.filename or '').lower()
        if name.endswith('.xls') or data[:4] == b'\xd0\xcf\x11\xe0':
            return self._read_xls(data)
        if name.endswith('.xlsx') or data[:2] == b'PK':
            return self._read_xlsx(data)
        raise UserError(_('Formato no soportado. Sube el reporte como .xls o .xlsx.'))

    def _read_xls(self, data):
        try:
            import xlrd
        except ImportError:
            raise UserError(_('Falta la librería xlrd para leer archivos .xls.'))
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)]
                for r in range(sheet.nrows)]

    def _read_xlsx(self, data):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_('Falta la librería openpyxl para leer archivos .xlsx.'))
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        return [['' if v is None else v for v in row]
                for row in sheet.iter_rows(values_only=True)]

    def _extract_period(self, matrix):
        head_text = ' '.join(
            c for row in matrix[:8] for c in row
            if isinstance(c, str) and c.strip())
        match = PERIOD_RE.search(head_text)
        parsed_start = parsed_end = None
        if match:
            parsed_start = _mk_date(match.group(1), match.group(2), match.group(3))
            parsed_end = _mk_date(match.group(4), match.group(5), match.group(6))

        start = self.week_start_override or parsed_start
        if not start:
            raise UserError(_('No se pudo leer el rango de fechas del reporte '
                              '("… al …"). Indica la fecha de la semana a mano.'))
        end = parsed_end or (start + timedelta(days=6))
        return start, end

    def _locate_columns(self, matrix, period_start):
        for idx, row in enumerate(matrix[:15]):
            norm = [_norm(c) for c in row]
            if 'id' not in norm or not any(n.startswith('nombre') for n in norm):
                continue
            cols = {
                'id': norm.index('id'),
                'name': next(i for i, n in enumerate(norm) if n.startswith('nombre')),
            }
            diff_i = next((i for i, n in enumerate(norm)
                           if 'dif' in n and '44' in n), None)
            total_i = next((i for i, n in enumerate(norm)
                            if 'total' in n and 'hora' in n), None)
            if diff_i is not None:
                cols['mode'], cols['value'] = 'diff', diff_i
            elif total_i is not None:
                cols['mode'], cols['value'] = 'total', total_i
            else:
                raise UserError(_("El encabezado no tiene ni 'Dif. vs 44h' ni "
                                  "'Total Horas'."))
            cols['total'] = total_i
            cols['day_blocks'] = self._locate_day_blocks(norm, period_start)
            return idx, cols
        raise UserError(_("No se encontró la fila de encabezados (columnas 'ID' y "
                          "'Nombre') en el archivo."))

    def _locate_day_blocks(self, norm_header, period_start):
        """Columnas (Entrada, Salida, Horas) y fecha de cada día del reporte."""
        blocks = []
        for i, cell in enumerate(norm_header):
            m = DAY_HEADER_RE.match(cell)
            if not m:
                continue
            day_date = _mk_date(m.group(2), m.group(3), period_start.year)
            if not day_date:
                day_date = period_start + timedelta(days=len(blocks))
            elif day_date < period_start:      # semana que cruza fin de año
                day_date = date(period_start.year + 1, day_date.month, day_date.day)
            blocks.append({'entrada': i, 'salida': i + 1,
                           'horas': i + 2, 'date': day_date})
        return blocks

    def _first_data_row(self, matrix, header_idx, cols):
        nxt = header_idx + 1
        if nxt < len(matrix):
            probe = matrix[nxt][cols['id']] if cols['id'] < len(matrix[nxt]) else None
            if not isinstance(probe, (int, float)):
                return nxt + 1
        return nxt

    # ------------------------------------------------------------------ #
    #  Resolución de empleado / tarifa
    # ------------------------------------------------------------------ #
    def _match_employee(self, name):
        Employee = self.env['hr.employee']
        name = (name or '').strip()
        if not name:
            return Employee
        exact = Employee.search([('name', '=ilike', name)], limit=2)
        if len(exact) == 1:
            return exact
        if len(exact) > 1:
            return Employee
        partial = Employee.search([('name', 'ilike', name)], limit=2)
        return partial if len(partial) == 1 else Employee

    def _resolve_hourly_rate(self, employee):
        salary = self.env['employee.salary'].search([
            ('employee_id', '=', employee.id),
            ('is_active', '=', True),
        ], order='effective_date desc', limit=1)
        return salary.base_salary_lps / MONTHLY_HOURS if salary else 0.0

    def _rrhh_employee(self, company):
        return self.env['hr.employee'].search([
            ('user_id', '=', self.env.uid),
            ('company_id', '=', company.id),
        ], limit=1)
